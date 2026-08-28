# -*- coding: utf-8 -*-
"""Dịch các slide bài giảng (ch1-ch5) sang tiếng Anh.

Đầu vào là bản pptx tiếng Việt đã hoàn thiện (đã qua build_decks.js /
build_ch5_practice.js, upgrade_decks.py và add_images.py) ở slides/0N-*.pptx —
tệp này giữ vai trò "nguồn" cho phiên bản tiếng Anh, giống cách lam_phu_de.py
dịch phụ đề từ videos/kich_ban_video.json chứ không sinh lại từ đầu.

Với mỗi đoạn văn (paragraph) trên slide, tra bảng dịch theo đúng nguyên văn
tiếng Việt rồi thay chữ ở run đầu tiên (giữ nguyên phông, cỡ, màu, in đậm...),
xoá chữ ở các run thừa. Số trang, mã mục (1.1, 7%...) giữ nguyên không dịch.
Ghi chú giảng bài (speaker notes) KHÔNG dịch — cô Hương giảng bằng tiếng Việt
nên notes vẫn để tiếng Việt ở cả bản EN của slide.

Chạy: python3 scripts/dich_slide_en.py
Thiếu bản dịch cho câu nào thì báo lỗi và dừng — không được bỏ sót.
"""
import os
import re
import sys
from pptx import Presentation

GIU_NGUYEN = re.compile(r"^[\d\.\-–%\s]+$")

# Chữ lặp lại ở mọi slide của mọi chương (chân trang, tên chương).
CHUNG = {
    "Je m'appelle Huong  •  GV. Đỗ Thùy Hương  •  EC1103":
        "Je m'appelle Huong  •  Instructor Đỗ Thùy Hương  •  EC1103",
    "GV. Đỗ Thùy Hương": "Instructor Đỗ Thùy Hương",
    "MỤC TIÊU": "OBJECTIVES",
    "NỘI DUNG": "AGENDA",
    "TỔNG KẾT": "SUMMARY",
    "ÔN TẬP & CHUẨN BỊ": "REVIEW & PREPARATION",
    "TÀI LIỆU": "REFERENCES",
    "HOẠT ĐỘNG NHÓM": "GROUP ACTIVITY",
    "NHIỆM VỤ CỦA NHÓM": "GROUP TASKS",
    "EC1103 – Kỹ năng giao tiếp và soạn thảo văn bản (2:1)  •  Lớp 261b, HK1 năm học 2026 – 2027":
        "EC1103 – Business Communication and Document Drafting Skills (2:1)  •  Class 261b, Semester 1, 2026 – 2027",
    "© Đỗ Thùy Hương, 2026 — Bài giảng biên soạn cho lớp giảng dạy trực tiếp. Vui lòng ghi nguồn khi sử dụng.":
        "© Đỗ Thùy Hương, 2026 — Lecture prepared for the in-person class. Please credit the source when reused.",
    "Sau chương này, sinh viên có thể": "After this chapter, students will be able to",
    "Chúng ta sẽ đi qua": "What we will cover",
    "Ba điều cần nhớ của chương": "Three things to remember from this chapter",
    "Câu hỏi ôn tập": "Review questions",
    "Tài liệu học tập": "Course materials",
    "Giáo trình chính": "Main textbook",
    "Hà Nam Khánh Giao (2023), Giáo trình Giao tiếp kinh doanh, NXB Tài chính.":
        "Hà Nam Khánh Giao (2023), Business Communication Textbook, Finance Publishing House.",
    "Tài liệu tham khảo": "Further reading",
    "Thái Trí Dũng (2012), Kỹ năng giao tiếp và thương lượng trong kinh doanh, NXB Lao động – Xã hội.  •  Nghị định 30/2020/NĐ-CP về công tác văn thư (dùng cho Chương 5 và phần thực hành).":
        "Thái Trí Dũng (2012), Communication and Negotiation Skills in Business, Labour – Social Affairs Publishing House.  •  Decree 30/2020/ND-CP on records and archival work (used for Chapter 5 and the practice sessions).",
    "Học liệu của giảng viên": "Instructor's own materials",
    "Slide bài giảng, tình huống và bài tập do GV. Đỗ Thùy Hương biên soạn; cung cấp sau mỗi buổi học trên nhóm lớp.":
        "Slides, case scenarios and exercises prepared by Instructor Đỗ Thùy Hương; shared after each session in the class group.",
}

for n in range(1, 6):
    CHUNG[f"Chương {n}"] = f"Chapter {n}"
    CHUNG[f"CHƯƠNG {n}"] = f"CHAPTER {n}"

CH1 = {
    "Tổng quan về giao tiếp": "An overview of communication",
    "trong kinh doanh": "in business",
    "Nền móng của mọi kỹ năng nghề nghiệp: hiểu đúng về giao tiếp trước khi luyện kỹ năng.":
        "The foundation of every professional skill: understand communication correctly before drilling techniques.",
    "LỘ TRÌNH HỌC PHẦN": "COURSE ROADMAP",
    "Chúng ta sẽ học cùng nhau thế nào?": "How will we learn together?",
    "Lý thuyết — 10 buổi": "Theory — 10 sessions",
    "Sáng T7 & CN  •  5 chương  •  phòng C0105": "Sat & Sun mornings  •  5 chapters  •  room C0105",
    "Thực hành — 3 bài": "Practice — 3 labs",
    "Chiều T7 / CN  •  phòng A0105 Mô phỏng Kinh tế": "Sat / Sun afternoons  •  room A0105 Economics Simulation",
    "Đánh giá — 3 cột điểm": "Grading — 3 components",
    "Chuyên cần  •  Quá trình  •  Thi cuối kỳ": "Attendance  •  Coursework  •  Final exam",
    "Trình bày  —  khái niệm, đặc điểm của giao tiếp trong kinh doanh và mô hình quá trình giao tiếp.":
        "Explain  —  the concept and features of business communication and the communication process model.",
    "Phân biệt  —  các phương tiện và hình thức giao tiếp; nhận diện ưu – nhược điểm của từng hình thức.":
        "Distinguish  —  communication channels and forms; identify the pros and cons of each.",
    "Phân tích  —  các yếu tố ảnh hưởng đến hiệu quả giao tiếp trong tình huống thực tế.":
        "Analyse  —  the factors affecting communication effectiveness in real situations.",
    "Vận dụng  —  các nguyên tắc giao tiếp để xử lý một tình huống giao tiếp kinh doanh cụ thể.":
        "Apply  —  communication principles to handle a specific business communication situation.",
    "1.1  —  Khái niệm, đặc điểm của giao tiếp trong kinh doanh": "1.1  —  Concept and features of business communication",
    "1.2  —  Các phương tiện giao tiếp": "1.2  —  Communication channels",
    "1.3  —  Các hình thức giao tiếp": "1.3  —  Forms of communication",
    "1.4  —  Các yếu tố ảnh hưởng đến quá trình giao tiếp": "1.4  —  Factors affecting the communication process",
    "1.5  —  Các nguyên tắc giao tiếp": "1.5  —  Principles of communication",
    "MỤC 1.1 – 1.2": "SECTIONS 1.1 – 1.2",
    "Bản chất và phương tiện giao tiếp": "The nature and channels of communication",
    "Hiểu đúng bản chất trước khi luyện kỹ năng: giao tiếp là gì, diễn ra qua những khâu nào, bằng phương tiện gì.":
        "Understand the nature first: what communication is, what stages it goes through, and by what channels.",
    "Giao tiếp và giao tiếp trong kinh doanh": "Communication and business communication",
    "Giao tiếp là gì?": "What is communication?",
    "Trao đổi thông tin để đạt một mục đích": "Exchanging information to achieve a purpose",
    "Trong kinh doanh?": "In business?",
    "Gắn với mục tiêu công việc, có ràng buộc": "Tied to work goals, with constraints",
    "Vì sao phải học?": "Why study it?",
    "Nhà tuyển dụng xếp vào nhóm đòi hỏi cao nhất": "Employers rank it among the most in-demand skills",
    "Đặc điểm của giao tiếp trong kinh doanh": "Features of business communication",
    "Luôn có mục đích": "Always purposeful",
    "Mỗi cuộc gặp phục vụ một mục tiêu công việc": "Every meeting serves a work objective",
    "Đa dạng chủ thể": "Diverse parties",
    "Mỗi đối tượng một chuẩn mực riêng": "Each party has its own norms",
    "Ràng buộc lợi ích – pháp lý": "Bound by interests and law",
    "Lời nói có thể tạo ra nghĩa vụ": "Words can create obligations",
    "Khoa học và nghệ thuật": "Both a science and an art",
    "Có nguyên tắc, nhưng cần linh hoạt": "Governed by principles, yet needs flexibility",
    "Mô hình quá trình giao tiếp": "The communication process model",
    "Giao tiếp là quá trình hai chiều — muốn sửa một cuộc giao tiếp thất bại, hãy dò lại từng khâu.":
        "Communication is a two-way process — to fix a failed exchange, trace back through each stage.",
    "Phương tiện giao tiếp: ngôn ngữ": "Communication channels: verbal",
    "Ngôn ngữ nói": "Spoken language",
    "Nhanh, giàu cảm xúc — nhưng lời nói gió bay": "Fast, expressive — but words fade away",
    "Ngôn ngữ viết": "Written language",
    "Chính xác, lưu được — nền tảng của Chương 5": "Precise, kept on record — the foundation of Chapter 5",
    "Nguyên tắc dùng từ": "Word-choice principles",
    "Rõ  •  lịch sự  •  tích cực  •  ngắn gọn": "Clear  •  polite  •  positive  •  concise",
    "Phương tiện giao tiếp: phi ngôn ngữ": "Communication channels: non-verbal",
    "Ánh mắt – nét mặt": "Eyes and facial expression",
    "Kênh biểu cảm mạnh nhất": "The most expressive channel",
    "Cử chỉ – tư thế": "Gestures and posture",
    "Thẳng và cởi mở tạo thiện cảm": "Upright and open postures build rapport",
    "Khoảng cách": "Distance",
    "Thân mật · cá nhân · xã giao · công cộng": "Intimate · personal · social · public",
    "Trang phục – giọng – giờ giấc": "Dress, voice and punctuality",
    "Đúng giờ cũng là một thông điệp": "Being on time is a message too",
    "TỪ NGỮ": "WORDS",
    "GIỌNG NÓI": "TONE OF VOICE",
    "CƠ THỂ": "BODY",
    "Mehrabian — với thông điệp cảm xúc": "Mehrabian — for emotional messages",
    "CON SỐ PHẢI NHỚ": "THE NUMBER TO REMEMBER",
    "Thông điệp cảm xúc được truyền đi thế nào?": "How is an emotional message actually conveyed?",
    "Nội dung lời nói — phần nhỏ nhất, dù ta thường đầu tư nhiều nhất vào đây":
        "The words themselves — the smallest share, even though we usually invest the most effort here",
    "Âm lượng, tốc độ, ngữ điệu, khoảng dừng": "Volume, pace, intonation, pauses",
    "NGÔN NGỮ CƠ THỂ": "BODY LANGUAGE",
    "Ánh mắt, nét mặt, cử chỉ, tư thế, khoảng cách": "Eyes, facial expression, gestures, posture, distance",
    "Nghiên cứu của Albert Mehrabian — chỉ áp dụng cho thông điệp mang tính CẢM XÚC, không phải mọi tình huống giao tiếp.":
        "Albert Mehrabian's study — applies only to EMOTIONAL messages, not to every communication situation.",
    "Các hình thức giao tiếp": "Forms of communication",
    "Bốn cặp này không loại trừ nhau — một cuộc giao tiếp nằm đâu đó trên cả bốn trục cùng lúc.":
        "These four pairs are not mutually exclusive — a single exchange sits somewhere on all four axes at once.",
    "MỤC 1.4 – 1.5": "SECTIONS 1.4 – 1.5",
    "Yếu tố ảnh hưởng và nguyên tắc": "Influencing factors and principles",
    "Điều gì làm hỏng một cuộc giao tiếp, và năm nguyên tắc giúp ta tránh được điều đó.":
        "What derails a conversation, and five principles that help us avoid it.",
    "Yếu tố ảnh hưởng đến quá trình giao tiếp": "Factors affecting the communication process",
    "Hỏng ở yếu tố nào thì sửa đúng yếu tố ấy, đừng đổ hết cho “nói chưa khéo”.":
        "Whichever factor breaks down, fix that exact factor — don't blame it all on \"poor phrasing\".",
    "Nguyên tắc giao tiếp trong kinh doanh": "Principles of business communication",
    "Tôn trọng  —  nhân cách, thời gian, lợi ích, khác biệt": "Respect  —  for character, time, interests and differences",
    "Thiện chí – hợp tác  —  thắng một cuộc cãi, thua một khách hàng": "Goodwill and cooperation  —  win an argument, lose a customer",
    "Lắng nghe trước  —  hiểu đúng rồi mới nói": "Listen first  —  understand correctly before speaking",
    "Phù hợp ngữ cảnh  —  đúng vai, đúng lúc, đúng kênh": "Fit the context  —  the right role, moment and channel",
    "Giữ chữ tín  —  đã hứa là làm": "Keep your word  —  a promise made is a promise kept",
    "Tình huống: buổi gặp đầu tiên thất bại": "Case: a failed first meeting",
    "TÌNH HUỐNG (thảo luận nhóm 4–5 sinh viên, 15 phút)": "SCENARIO (groups of 4–5 students, 15 minutes)",
    "Nhân viên kinh doanh A đến gặp khách hàng lần đầu: đến trễ 10 phút vì kẹt xe nhưng không báo trước; mặc áo thun vì “cuối tuần”; vừa ngồi đã mở máy giới thiệu sản phẩm liên tục 20 phút; điện thoại đổ chuông 2 lần và A đều bắt máy. Kết thúc buổi gặp, khách hàng nói “để anh xem lại rồi báo em sau” và không phản hồi nữa.":
        "Salesperson A meets a client for the first time: arrives 10 minutes late because of traffic, without notice; wears a T-shirt because \"it's the weekend\"; sits down and pitches the product non-stop for 20 minutes; the phone rings twice and A answers both times. The meeting ends with the client saying \"let me think it over and get back to you\" — and never replying again.",
    "NHIỆM VỤ CỦA NHÓM": "GROUP TASKS",
    "Liệt kê tất cả các lỗi giao tiếp của A và xếp mỗi lỗi vào một khâu trong mô hình quá trình giao tiếp.":
        "List all of A's communication mistakes and map each one to a stage in the communication process model.",
    "Mỗi lỗi vi phạm nguyên tắc giao tiếp nào ở mục 1.5?": "Which principle from section 1.5 does each mistake violate?",
    "Xây dựng “kịch bản chuẩn” 5 bước cho buổi gặp khách hàng đầu tiên và cử đại diện trình bày trước lớp (3 phút).":
        "Build a 5-step \"standard script\" for a first client meeting and have a representative present it to the class (3 minutes).",
    "Có mục đích và có luật chơi": "Purposeful, and governed by rules",
    "Không phải trò chuyện ngẫu nhiên": "Not just random chit-chat",
    "Hỏng ở khâu nào, dò lại khâu đó": "Trace the failure back to its exact stage",
    "Năm khâu, cộng thêm nhiễu": "Five stages, plus noise",
    "Phi ngôn ngữ mạnh hơn ta nghĩ": "Non-verbal cues are stronger than we think",
    "Nó nói trước, và nói to hơn lời": "It speaks first, and louder than words",
    "Phân tích các thành phần của mô hình quá trình giao tiếp qua một ví dụ thực tế của chính bạn.":
        "Analyse the components of the communication process model using one of your own real-life examples.",
    "So sánh ưu – nhược điểm của giao tiếp bằng lời nói và bằng văn bản trong kinh doanh.":
        "Compare the pros and cons of verbal and written communication in business.",
    "Vì sao nói giao tiếp kinh doanh “vừa là khoa học, vừa là nghệ thuật”?":
        "Why is business communication said to be \"both a science and an art\"?",
    "Nêu và minh họa 5 nguyên tắc giao tiếp trong kinh doanh.": "State and illustrate the 5 principles of business communication.",
    "CHUẨN BỊ CHO BUỔI SAU:  Chương 2 – Kỹ năng giao tiếp chuyên nghiệp. Mỗi nhóm chuẩn bị một bài thuyết trình 3 phút về chủ đề tự chọn để thực hành trên lớp.":
        "PREPARE FOR NEXT SESSION:  Chapter 2 – Professional communication skills. Each group prepares a 3-minute presentation on a topic of their choice to practise in class.",
}

CH2 = {
    "Các kỹ năng giao tiếp": "Professional communication",
    "chuyên nghiệp": "skills",
    "Bốn kỹ năng dùng hằng ngày suốt sự nghiệp: gây ấn tượng, thuyết trình, lắng nghe và điện thoại.":
        "Four skills used every day of your career: making an impression, presenting, listening, and phone calls.",
    "Tạo dựng  —  ấn tượng ban đầu chuyên nghiệp và thực hiện đúng các nghi thức xã giao công sở.":
        "Build  —  a professional first impression and follow proper workplace etiquette.",
    "Chuẩn bị và trình bày  —  một bài thuyết trình có cấu trúc, tự tin trước đám đông.":
        "Prepare and deliver  —  a well-structured presentation, confidently in front of an audience.",
    "Thực hành  —  lắng nghe chủ động và đặt câu hỏi hiệu quả trong hội thoại công việc.":
        "Practise  —  active listening and effective questioning in workplace conversations.",
    "Giao tiếp qua điện thoại  —  đúng chuẩn mực nghề nghiệp ở cả vai gọi đi và nghe máy.":
        "Communicate by phone  —  to a professional standard both making and receiving calls.",
    "2.1  —  Kỹ năng tạo ấn tượng ban đầu và xã giao": "2.1  —  Making a first impression and etiquette",
    "2.2  —  Kỹ năng thuyết trình": "2.2  —  Presentation skills",
    "2.3  —  Kỹ năng lắng nghe và đặt câu hỏi": "2.3  —  Listening and questioning skills",
    "2.4  —  Kỹ năng giao tiếp qua điện thoại": "2.4  —  Telephone communication skills",
    "MỤC 2.1": "SECTION 2.1",
    "Ấn tượng ban đầu và nghi thức xã giao": "First impressions and etiquette",
    "Hai mươi giây đầu tiên quyết định phần lớn cách người khác nhìn nhận bạn.":
        "The first twenty seconds decide most of how others perceive you.",
    "Ấn tượng ban đầu — chỉ có một lần": "A first impression — you only get one",
    "CON SỐ PHẢI NHỚ": "THE NUMBER TO REMEMBER",
    "Quy tắc 4 × 20 — bốn cửa ải của ấn tượng ban đầu": "The 4 × 20 rule — four gates of a first impression",
    "GIÂY": "SECONDS",
    "Đối phương hình thành đánh giá tổng thể": "The other person forms an overall judgement",
    "BƯỚC CHÂN": "STEPS",
    "Dáng đi, tư thế được đọc từ xa": "Gait and posture are read from a distance",
    "CENTIMET": "CENTIMETRES",
    "Ánh mắt và nụ cười trên gương mặt": "The eyes and the smile on your face",
    "TỪ": "WORDS",
    "Lời chào và giới thiệu đầu tiên": "The first greeting and introduction",
    "Ấn tượng ban đầu hình thành gần như tức thì và rất khó đảo ngược — nó phải được CHUẨN BỊ, không phó mặc cho may mắn.":
        "A first impression forms almost instantly and is very hard to reverse — it must be PREPARED, not left to chance.",
    "Nghi thức xã giao cơ bản": "Basic social etiquette",
    "Chào hỏi và giới thiệu": "Greeting and introductions",
    "Người nhỏ chào người lớn, nhân viên chào cấp trên trước; giới thiệu người ít quan trọng với người quan trọng hơn; nói rõ tên – chức danh – đơn vị.":
        "The younger greets the elder, staff greet superiors first; introduce the less senior person to the more senior one; state name, title and unit clearly.",
    "Bắt tay": "Handshake",
    "Đứng dậy, nhìn vào mắt, siết vừa phải 2–3 giây; người có vị thế cao hơn hoặc phụ nữ chủ động đưa tay trước; không bắt quá lỏng, quá chặt, quá lâu.":
        "Stand up, make eye contact, hold for a moderate 2–3 seconds; the higher-status person or a woman offers their hand first; don't grip too loosely, too firmly, or too long.",
    "Trao – nhận danh thiếp": "Giving and receiving business cards",
    "Trao bằng hai tay, mặt chữ hướng về người nhận; nhận bằng hai tay, đọc qua rồi mới cất — đừng nhét ngay vào túi quần.":
        "Present with both hands, text facing the recipient; receive with both hands, read it before putting it away — don't stuff it straight into a pocket.",
    "Ứng xử không gian chung": "Behaviour in shared spaces",
    "Thang máy, phòng họp, bàn làm việc mở: giữ trật tự, nhường lối, gõ cửa trước khi vào, không nói to chuyện riêng.":
        "Elevators, meeting rooms, open-plan desks: keep order, give way, knock before entering, don't talk loudly about personal matters.",
    "MỤC 2.2": "SECTION 2.2",
    "Kỹ năng thuyết trình": "Presentation skills",
    "Từ chuẩn bị đến trình bày: làm sao để người nghe nhớ được điều bạn muốn nói.":
        "From preparation to delivery: how to make the audience remember what you want to say.",
    "Thuyết trình: 5 bước chuẩn bị": "Presenting: 5 steps of preparation",
    "PHÂN TÍCH NGƯỜI NGHE": "AUDIENCE",
    "XÁC ĐỊNH MỤC TIÊU": "OBJECTIVE",
    "XÂY DỰNG NỘI DUNG": "CONTENT",
    "THIẾT KẾ SLIDE": "DESIGN",
    "LUYỆN TẬP": "REHEARSE",
    "Họ là ai, quan tâm gì, mong đợi điều gì?": "Who are they, what do they care about, what do they expect?",
    "Sau bài nói, người nghe biết / tin / làm gì?": "After the talk, what should listeners know / believe / do?",
    "Chọn 3 ý chính, dẫn chứng, ví dụ.": "Pick 3 main points, evidence, examples.",
    "Ít chữ, nhiều hình; slide hỗ trợ chứ không thay người nói.": "Few words, more visuals; slides support the speaker, they don't replace them.",
    "Tập nói to, canh giờ, dự phòng câu hỏi khó.": "Practise speaking aloud, time yourself, prepare for tough questions.",
    "Kinh nghiệm: 1 phút thuyết trình cần khoảng 1 giờ chuẩn bị nếu chủ đề mới — thời lượng luyện tập là thứ khán giả “nhìn thấy” rõ nhất.":
        "Rule of thumb: 1 minute of presenting needs about 1 hour of preparation for a new topic — rehearsal time is what the audience \"sees\" most clearly.",
    "Cấu trúc bài thuyết trình": "Presentation structure",
    "Mở đầu (10–15%) — giành lấy sự chú ý": "Opening (10–15%) — capture attention",
    "Câu hỏi gây tò mò, con số ấn tượng, câu chuyện ngắn; giới thiệu bản thân và cho người nghe biết lộ trình bài nói.":
        "An intriguing question, a striking figure, a short story; introduce yourself and outline the talk's roadmap.",
    "Thân bài (70–80%) — tối đa 3 ý chính": "Body (70–80%) — at most 3 main points",
    "Mỗi ý: luận điểm → dẫn chứng → ví dụ; có câu chuyển ý rõ ràng; đặt ý mạnh nhất ở đầu hoặc cuối.":
        "For each point: claim → evidence → example; use clear transitions; place your strongest point first or last.",
    "Kết luận (10–15%) — đọng lại một điều": "Conclusion (10–15%) — leave one thing behind",
    "Tóm tắt 3 ý, nhấn thông điệp cốt lõi, kêu gọi hành động cụ thể; chuẩn bị sẵn phần hỏi – đáp.":
        "Summarise the 3 points, stress the core message, call for a specific action; be ready for Q&A.",
    "Trình bày tự tin trước đám đông": "Presenting confidently to an audience",
    "Ngôn ngữ cơ thể": "Body language",
    "Đứng vững, mở vai, di chuyển có chủ đích; giao tiếp mắt luân phiên khắp phòng; tay minh họa tự nhiên, không đút túi hay vặn bút.":
        "Stand firm, open shoulders, move with purpose; make eye contact around the whole room in turn; use natural hand gestures, don't pocket your hands or fidget with a pen.",
    "Giọng nói": "Voice",
    "Nói to rõ, thay đổi tốc độ và ngữ điệu; dừng 1–2 giây trước ý quan trọng — khoảng lặng là gia vị của bài nói.":
        "Speak loudly and clearly, vary pace and intonation; pause 1–2 seconds before a key point — silence is the seasoning of a talk.",
    "Vượt qua run sợ": "Overcoming nerves",
    "Run là bình thường; chuẩn bị kỹ + đến sớm làm quen không gian + hít thở sâu + tập trung vào thông điệp thay vì bản thân.":
        "Nerves are normal; prepare thoroughly + arrive early to get used to the room + breathe deeply + focus on the message rather than yourself.",
    "Xử lý câu hỏi": "Handling questions",
    "Lắng nghe hết câu hỏi, cảm ơn, trả lời ngắn gọn; chưa chắc thì hẹn trả lời sau — trung thực hơn là đoán bừa.":
        "Hear the whole question out, thank the asker, answer concisely; if unsure, promise to follow up later — honesty beats guessing.",
    "MỤC 2.3 – 2.4": "SECTIONS 2.3 – 2.4",
    "Lắng nghe, đặt câu hỏi và giao tiếp qua điện thoại": "Listening, questioning and telephone communication",
    "Hai kỹ năng ít được dạy nhất nhưng quyết định nhất trong công việc hằng ngày.":
        "The two least-taught skills, yet the most decisive ones in everyday work.",
    "Nghe khác lắng nghe: 5 mức độ": "Hearing vs. listening: 5 levels",
    "Kỹ năng đặt câu hỏi": "Questioning skills",
    "Câu hỏi đóng": "Closed questions",
    "Trả lời Có/Không hoặc một dữ kiện: “Anh đã nhận được báo giá chưa?” — dùng để xác nhận, chốt thông tin.":
        "Answered with Yes/No or one fact: \"Have you received the quotation yet?\" — used to confirm and pin down information.",
    "Câu hỏi mở": "Open questions",
    "Bắt đầu bằng Vì sao / Như thế nào / Điều gì: khơi người đối diện chia sẻ — dùng để khám phá nhu cầu.":
        "Start with Why / How / What: invite the other person to open up — used to explore needs.",
    "Câu hỏi thăm dò – đào sâu": "Probing questions",
    "“Cụ thể là…?”, “Anh có thể cho ví dụ?” — làm rõ chi tiết sau câu trả lời chung chung.":
        "\"Specifically…?\", \"Could you give an example?\" — clarify detail after a vague answer.",
    "Lưu ý khi hỏi": "Notes on asking questions",
    "Mỗi lần một câu; tránh câu hỏi dồn ép, mớm cung; hỏi xong thì im lặng chờ — đừng tự trả lời thay.":
        "One question at a time; avoid pressuring or leading questions; after asking, wait in silence — don't answer for them.",
    "Điện thoại: cuộc gọi đi chuyên nghiệp": "Phone calls: making a professional call",
    "Chuẩn bị trước khi gọi  —  mục đích, nội dung chính, giấy bút ghi chú; chọn thời điểm phù hợp (tránh sớm quá, muộn quá, giờ nghỉ trưa).":
        "Prepare before calling  —  purpose, key points, pen and paper for notes; pick a suitable time (avoid too early, too late, or the lunch break).",
    "Mở đầu đúng nghi thức  —  chào, xưng danh và đơn vị, xin phép: “Em chào anh, em là… từ công ty… Anh có tiện nghe máy 5 phút không ạ?”":
        "Open with proper etiquette  —  greet, state your name and organisation, ask permission: \"Hello, this is … from … company. Do you have 5 minutes to talk?\"",
    "Trình bày gọn, kiểm tra hiểu  —  đi thẳng vào việc, nói rõ ràng; tóm tắt lại thỏa thuận: thời gian, địa điểm, việc cần làm.":
        "Be concise, check understanding  —  get straight to the point, speak clearly; recap the agreement: time, place, and next steps.",
    "Kết thúc lịch sự  —  cảm ơn, chào; để người có vị thế cao hơn / khách hàng gác máy trước.":
        "End politely  —  thank them, say goodbye; let the more senior person / the customer hang up first.",
    "Nghe máy và văn hóa điện thoại nơi làm việc": "Answering calls and phone etiquette at work",
    "Khi nghe máy": "When answering",
    "Nhấc máy trong ~3 hồi chuông; chào và xưng danh đơn vị; giọng tươi — người gọi “nghe thấy” nụ cười của bạn; ghi chú và nhắc lại lời nhắn.":
        "Pick up within ~3 rings; greet and state your organisation; sound upbeat — the caller can \"hear\" your smile; take notes and repeat back the message.",
    "Khi người cần gặp vắng mặt": "When the person they want is unavailable",
    "Đề nghị để lại lời nhắn: ai gọi – việc gì – số liên lạc – hẹn phản hồi; chuyển lời nhắn đầy đủ, đúng hạn.":
        "Offer to take a message: who called – what it's about – contact number – when to expect a reply; pass the message on fully and on time.",
    "Di động nơi công sở": "Mobile phones at the office",
    "Chế độ im lặng trong cuộc họp; không nghe điện riêng khi đang tiếp khách; nhắn tin/Zalo công việc cũng cần đúng chuẩn mực như email.":
        "Silent mode during meetings; don't take personal calls while hosting a guest; work texts/chat messages need the same standard as email.",
    "Thực hành 2 kỹ năng ngay tại lớp": "Practising two skills right in class",
    "NHIỆM VỤ KÉP (nhóm 4–5 sinh viên, 25 phút chuẩn bị + trình diễn)": "DUAL TASK (groups of 4–5 students, 25 minutes to prepare + perform)",
    "Nhóm bốc thăm một tình huống: (a) gọi điện lần đầu cho khách hàng tiềm năng để hẹn gặp giới thiệu sản phẩm; hoặc (b) gọi điện xử lý việc giao hàng trễ cho khách đang khó chịu.":
        "Each group draws a scenario: (a) a first call to a prospective client to arrange a product demo; or (b) a call handling a late delivery for an upset customer.",
    "Đóng vai cuộc gọi 3 phút trước lớp: một bạn vai nhân viên, một bạn vai khách hàng; cả lớp chấm theo checklist mục 2.4.":
        "Role-play a 3-minute call in front of the class: one student as the employee, one as the customer; the class scores it against the section 2.4 checklist.",
    "Nhóm còn lại thuyết trình 3 phút chủ đề đã chuẩn bị ở nhà; cả lớp nhận xét theo cấu trúc mở – thân – kết và ngôn ngữ cơ thể.":
        "The other groups give a 3-minute presentation on a topic prepared at home; the class comments on the opening–body–conclusion structure and body language.",
    "Mỗi nhóm rút ra 3 điều sẽ làm khác đi nếu được thực hiện lại.": "Each group identifies 3 things they would do differently next time.",
    "Ấn tượng ban đầu được chuẩn bị, không phải may mắn": "A first impression is prepared, not left to luck",
    "Trang phục – thần thái – lời chào – nghi thức xã giao: tất cả đều luyện được trước.":
        "Attire, demeanour, greeting, etiquette: all of it can be rehearsed beforehand.",
    "Thuyết trình hay bắt đầu từ người nghe": "A good presentation starts with the audience",
    "Phân tích khán giả → mục tiêu → 3 ý chính → luyện tập; nói với người nghe, không nói với slide.":
        "Analyse the audience → set the objective → 3 main points → rehearse; speak to your audience, not to your slides.",
    "Lắng nghe và đặt câu hỏi là kỹ năng “bán hàng” giỏi nhất": "Listening and questioning is the best \"selling\" skill",
    "Hiểu đúng nhu cầu trước, trình bày sau; điện thoại chuyên nghiệp là bộ mặt âm thanh của doanh nghiệp.":
        "Understand needs correctly before presenting; a professional phone manner is the audible face of a business.",
    "Trình bày quy tắc 4×20 và cách vận dụng trong buổi phỏng vấn xin việc.": "Explain the 4×20 rule and how to apply it in a job interview.",
    "Xây dựng cấu trúc chi tiết cho bài thuyết trình 5 phút giới thiệu một sản phẩm tự chọn.":
        "Build a detailed structure for a 5-minute presentation introducing a product of your choice.",
    "Phân biệt 5 mức độ lắng nghe; cho ví dụ về lắng nghe thấu cảm trong công việc.":
        "Distinguish the 5 levels of listening; give an example of empathetic listening at work.",
    "Soạn kịch bản cuộc gọi hẹn gặp khách hàng theo 4 bước chuẩn.": "Draft a script for a client call following the 4 standard steps.",
    "CHUẨN BỊ CHO BUỔI SAU:  Chương 3 – Giao tiếp trong các tình huống đặc thù. Mỗi nhóm sưu tầm một tình huống giao tiếp khó xử có thật tại nơi làm việc (giữ ẩn danh) để thảo luận.":
        "PREPARE FOR NEXT SESSION:  Chapter 3 – Communication in specific settings. Each group collects one real, anonymised workplace communication dilemma to discuss.",
}

CH3 = {
    "Giao tiếp trong các": "Communication in",
    "tình huống đặc thù": "specific settings",
    "Cùng một kỹ năng, mỗi bối cảnh một luật chơi: nội bộ, khách hàng, bàn tiệc và đa văn hóa.":
        "The same skill, different rules for each setting: internal, customers, banquets and cross-cultural work.",
    "Ứng xử phù hợp  —  với cấp trên, cấp dưới và đồng nghiệp trong môi trường nội bộ tổ chức.":
        "Behave appropriately  —  with superiors, subordinates and colleagues within the organisation.",
    "Giao tiếp chuyên nghiệp  —  với khách hàng, đối tác, cơ quan nhà nước và truyền thông; xử lý được phàn nàn của khách.":
        "Communicate professionally  —  with customers, partners, state agencies and the press; handle customer complaints.",
    "Thực hiện đúng  —  nghi thức giao tiếp trên bàn tiệc trong hoạt động kinh doanh.":
        "Follow correctly  —  banquet etiquette in business activities.",
    "Thích ứng  —  với khác biệt văn hóa khi làm việc trong môi trường đa văn hóa.":
        "Adapt  —  to cultural differences when working in a cross-cultural environment.",
    "3.1  —  Giao tiếp trong môi trường nội bộ tổ chức": "3.1  —  Communication within the organisation",
    "3.2  —  Giao tiếp với khách hàng, đối tác, cơ quan nhà nước và truyền thông":
        "3.2  —  Communication with customers, partners, state agencies and the press",
    "3.3  —  Giao tiếp trên bàn tiệc": "3.3  —  Banquet communication",
    "3.4  —  Giao tiếp trong môi trường đa văn hóa": "3.4  —  Cross-cultural communication",
    "MỤC 3.1": "SECTION 3.1",
    "Giao tiếp trong nội bộ tổ chức": "Communication within the organisation",
    "Với cấp trên, cấp dưới và đồng nghiệp — mỗi mối quan hệ một cách ứng xử.":
        "With superiors, subordinates and colleagues — each relationship has its own approach.",
    "Giao tiếp với cấp trên": "Communicating with superiors",
    "Khi nhận nhiệm vụ": "When receiving a task",
    "Lắng nghe – ghi chú – hỏi lại cho rõ yêu cầu, thời hạn, nguồn lực; xác nhận lại bằng tin nhắn/email để hai bên cùng hiểu một cách.":
        "Listen – take notes – ask for clarity on requirements, deadlines and resources; confirm by message/email so both sides share the same understanding.",
    "Khi báo cáo": "When reporting",
    "Chủ động, đúng hạn, kết quả trước – diễn giải sau; báo tin xấu sớm kèm phương án xử lý, không che giấu.":
        "Be proactive and on time, results first, explanation after; report bad news early with a proposed fix, never hide it.",
    "Khi có ý kiến khác": "When you disagree",
    "Chọn đúng lúc, đúng chỗ (thường là riêng tư); trình bày trên cơ sở dữ liệu và lợi ích chung; tôn trọng quyết định cuối cùng.":
        "Choose the right time and place (usually private); argue from data and shared interests; respect the final decision.",
    "Giao tiếp với cấp dưới và đồng nghiệp": "Communicating with subordinates and colleagues",
    "Giao việc": "Assigning work",
    "Rõ mục tiêu – thời hạn – tiêu chuẩn; giao việc kèm niềm tin và nguồn lực; kiểm tra tiến độ đúng mức, không quản lý vụn vặt.":
        "Be clear on the goal, deadline and standard; delegate with trust and resources; check progress at a sensible cadence, don't micromanage.",
    "Khen và phê bình": "Praise and criticism",
    "Khen công khai, kịp thời, cụ thể; phê bình riêng tư, nhắm vào hành vi chứ không nhắm vào con người, kèm hướng khắc phục.":
        "Praise publicly, promptly and specifically; criticise privately, target the behaviour not the person, and offer a way to fix it.",
    "Với đồng nghiệp": "With colleagues",
    "Tôn trọng, hợp tác, chia sẻ thông tin; tranh luận về công việc chứ không công kích cá nhân; tránh bè phái, tán chuyện sau lưng.":
        "Be respectful, cooperative, share information; debate the work, not the person; avoid cliques and gossiping behind people's backs.",
    "Họp hiệu quả": "Effective meetings",
    "Có chương trình gửi trước; đến đúng giờ; phát biểu ngắn gọn vào trọng tâm; có kết luận, biên bản và người chịu trách nhiệm từng việc.":
        "Send the agenda in advance; arrive on time; keep remarks brief and on point; end with conclusions, minutes and an owner for each action.",
    "MỤC 3.2": "SECTION 3.2",
    "Giao tiếp với bên ngoài tổ chức": "Communicating outside the organisation",
    "Khách hàng, đối tác, cơ quan nhà nước và truyền thông — bốn nhóm, bốn luật chơi.":
        "Customers, partners, state agencies and the press — four groups, four sets of rules.",
    "Giao tiếp với khách hàng": "Communicating with customers",
    "Tâm thế phục vụ": "A service mindset",
    "Khách hàng nuôi sống doanh nghiệp; mỗi điểm tiếp xúc (chào đón, tư vấn, giao hàng, hậu mãi) đều là khoảnh khắc xây hoặc phá niềm tin.":
        "Customers sustain the business; every touchpoint (welcome, advice, delivery, after-sales) either builds or breaks trust.",
    "Nguyên tắc vàng": "Golden rules",
    "Chào đón niềm nở – gọi tên khách khi có thể; lắng nghe nhu cầu trước khi giới thiệu; nói sự thật về sản phẩm; giữ lời hứa về thời hạn.":
        "Greet warmly – use the customer's name when possible; listen to their needs before pitching; tell the truth about the product; keep promised deadlines.",
    "Điều tối kỵ": "Absolute no-nos",
    "Tranh cãi thắng – thua với khách; hứa quá khả năng; đổ lỗi cho đồng nghiệp, cho quy trình; bỏ mặc khách sau khi bán xong.":
        "Arguing to \"win\" against a customer; over-promising; blaming colleagues or processes; abandoning the customer once the sale is made.",
    "Xử lý phàn nàn của khách hàng — quy trình LAST": "Handling customer complaints — the LAST process",
    "Với đối tác, cơ quan nhà nước và truyền thông": "With partners, state agencies and the press",
    "Đối tác kinh doanh": "Business partners",
    "Bình đẳng, giữ chữ tín, minh bạch thông tin; quan hệ lâu dài quan trọng hơn lợi thế ngắn hạn — nền tảng cho đàm phán ở Chương 4.":
        "Be equal, keep your word, be transparent; a long-term relationship matters more than a short-term edge — the foundation for negotiation in Chapter 4.",
    "Cơ quan nhà nước": "State agencies",
    "Đúng thủ tục, đúng thẩm quyền, hồ sơ – văn bản chuẩn thể thức (Chương 5); tác phong nghiêm túc, đúng hẹn; tuyệt đối không “đi tắt” trái quy định.":
        "Follow proper procedure and authority, correctly formatted dossiers and documents (Chapter 5); be serious and punctual; never cut corners against regulations.",
    "Truyền thông – báo chí": "Media and the press",
    "Chỉ người được ủy quyền phát ngôn; thông tin nhất quán, trung thực; khi có khủng hoảng: phản hồi nhanh, nhận trách nhiệm đúng phần của mình, không né tránh.":
        "Only an authorised spokesperson speaks; keep information consistent and truthful; in a crisis: respond quickly, own your share of responsibility, don't dodge.",
    "MỤC 3.3 – 3.4": "SECTIONS 3.3 – 3.4",
    "Bàn tiệc và môi trường đa văn hóa": "Banquets and cross-cultural settings",
    "Nơi công việc vẫn tiếp diễn dù không ai nhắc đến công việc.": "Where work carries on even though nobody mentions work.",
    "Giao tiếp trên bàn tiệc": "Banquet communication",
    "Trước bữa tiệc": "Before the meal",
    "Xác nhận tham dự đúng hạn; đến đúng giờ; trang phục theo tính chất tiệc; chờ chủ tiệc mời và xếp chỗ — vị trí ngồi thể hiện thứ bậc.":
        "RSVP on time; arrive punctually; dress to match the occasion; wait for the host to invite you in and seat you — seating reflects hierarchy.",
    "Trong bữa ăn": "During the meal",
    "Chủ tiệc bắt đầu trước; dùng dụng cụ từ ngoài vào trong; không nói khi đang nhai, không gõ đũa, không xoay đĩa thức ăn về phía mình liên tục.":
        "The host starts first; use cutlery from the outside in; don't talk with your mouth full, don't tap chopsticks, don't keep spinning the lazy Susan toward yourself.",
    "Chúc rượu – cụng ly": "Toasting and clinking glasses",
    "Người vị thế thấp nâng ly thấp hơn khi cụng; chúc ngắn gọn, đúng đối tượng; tôn trọng người không dùng rượu bia — không ép.":
        "The more junior person clinks with their glass held lower; keep toasts short and address the right person; respect those who don't drink alcohol — never pressure them.",
    "Câu chuyện trên bàn tiệc": "Table conversation",
    "Chủ đề nhẹ nhàng: ẩm thực, thể thao, du lịch, quê quán; tránh chính trị, tôn giáo, thu nhập, đời tư; công việc chỉ bàn khi chủ tiệc gợi mở.":
        "Keep it light: food, sport, travel, hometowns; avoid politics, religion, income, private life; only discuss work if the host brings it up.",
    "Giao tiếp đa văn hóa: nhận diện khác biệt": "Cross-cultural communication: spotting the differences",
    "Cách nói": "Speaking style",
    "Văn hóa “nói thẳng” (Đức, Mỹ, Hà Lan) đánh giá cao sự rõ ràng; văn hóa “nói vòng” (Nhật, Hàn, Việt) ưu tiên giữ thể diện — “để chúng tôi xem xét” có thể là lời từ chối.":
        "\"Direct\" cultures (Germany, the US, the Netherlands) value clarity; \"indirect\" cultures (Japan, Korea, Vietnam) prioritise saving face — \"we'll consider it\" can mean no.",
    "Thứ bậc và ra quyết định": "Hierarchy and decision-making",
    "Nơi coi trọng tôn ti (Nhật, Hàn, Trung): đúng vai, đúng cấp, quyết định tập thể chậm mà chắc; nơi bình đẳng (Bắc Âu, Úc): gọi tên, tranh luận thẳng với sếp là bình thường.":
        "Where hierarchy matters (Japan, Korea, China): stay in your role and rank, group decisions are slow but solid; where it's egalitarian (Nordics, Australia): first names and openly debating the boss are normal.",
    "Thời gian và cam kết": "Time and commitments",
    "Văn hóa giờ giấc chặt (Đức, Nhật, Thụy Sĩ): trễ 5 phút là thất lễ; văn hóa thời gian linh hoạt: quan hệ đi trước, tiến độ đi sau — cần chốt mốc bằng văn bản.":
        "Strict-time cultures (Germany, Japan, Switzerland): 5 minutes late is disrespectful; flexible-time cultures: relationships come before schedules — pin down milestones in writing.",
    "Cử chỉ và kiêng kỵ": "Gestures and taboos",
    "Cùng một cử chỉ mang nghĩa khác nhau giữa các nước; màu sắc, con số, quà tặng đều có thể nhạy cảm — tra cứu trước khi gặp đối tác nước ngoài.":
        "The same gesture can mean different things in different countries; colours, numbers and gifts can all be sensitive — look them up before meeting a foreign partner.",
    "Nguyên tắc thích ứng đa văn hóa": "Principles for cross-cultural adaptation",
    "Tìm hiểu trước  —  văn hóa giao tiếp, nghi thức chào hỏi, kiêng kỵ của đối tác trước mỗi cuộc gặp quan trọng.":
        "Research beforehand  —  your partner's communication style, greeting customs and taboos before any important meeting.",
    "Quan sát và điều chỉnh  —  để ý cách đối tác chào, trao danh thiếp, giữ khoảng cách… và ứng xử tương thích.":
        "Observe and adjust  —  notice how they greet, exchange cards, keep distance… and match your behaviour accordingly.",
    "Không suy diễn theo chuẩn của mình  —  một hành vi “kỳ lạ” có thể hoàn toàn bình thường trong văn hóa của họ — hỏi lịch sự thay vì phán xét.":
        "Don't judge by your own standards  —  a \"strange\" behaviour may be perfectly normal in their culture — ask politely instead of judging.",
    "Nói chậm, rõ, xác nhận lại bằng văn bản  —  khi khác ngôn ngữ: tránh tiếng lóng, thành ngữ; tóm tắt thỏa thuận qua email sau cuộc họp.":
        "Speak slowly and clearly, confirm in writing  —  across languages: avoid slang and idioms; summarise the agreement by email after the meeting.",
    "Khiêm tốn và cầu thị  —  sẵn sàng xin lỗi khi lỡ phạm điều kiêng kỵ; thiện chí học hỏi luôn được ghi nhận ở mọi nền văn hóa.":
        "Stay humble and eager to learn  —  be ready to apologise if you break a taboo; a genuine willingness to learn is valued in every culture.",
    "Hai tình huống khó — xử lý ngay tại lớp": "Two tough situations — handled right in class",
    "TÌNH HUỐNG (nhóm 4–5 sinh viên, 20 phút, bốc thăm 1 trong 2)": "SCENARIO (groups of 4–5 students, 20 minutes, draw one of two)",
    "Tình huống A: Khách hàng đến quầy lớn tiếng vì sản phẩm lỗi lần thứ hai trong tháng, nhiều khách khác đang nhìn. Tình huống B: Công ty tiếp đoàn đối tác Nhật Bản lần đầu — nhóm được giao chuẩn bị kịch bản đón tiếp và một bữa tiệc tối.":
        "Scenario A: A customer comes to the counter shouting about a faulty product for the second time this month, with other customers watching. Scenario B: The company is hosting a Japanese partner delegation for the first time — the group must prepare the welcome plan and a dinner banquet.",
    "Tình huống A: viết kịch bản xử lý theo đúng 4 bước LAST và đóng vai trước lớp (nhân viên – khách hàng – quản lý).":
        "Scenario A: write a response script following the 4 LAST steps and role-play it in front of the class (employee – customer – manager).",
    "Tình huống B: lập danh sách những việc phải làm và những điều tuyệt đối tránh (chào hỏi, danh thiếp, chỗ ngồi, quà tặng, chủ đề trò chuyện).":
        "Scenario B: list the things to do and the things to absolutely avoid (greetings, business cards, seating, gifts, conversation topics).",
    "Cả lớp nhận xét chéo: điều gì đã đúng chuẩn mực của chương, điều gì cần điều chỉnh?":
        "The class cross-reviews: what matched the chapter's standards, and what needs adjusting?",
    "Nội bộ vững thì đối ngoại mới mạnh": "A strong inside makes for a strong outside",
    "Nhận việc – báo cáo – phản hồi với cấp trên; giao việc – khen chê với cấp dưới: đều có chuẩn mực học được.":
        "Receiving tasks, reporting and giving feedback to superiors; delegating and praising/criticising subordinates — all follow learnable standards.",
    "Khách hàng phàn nàn là cơ hội": "A customer complaint is an opportunity",
    "LAST: Lắng nghe – Xin lỗi – Giải quyết – Cảm ơn; đừng thắng cuộc cãi để rồi mất khách hàng.":
        "LAST: Listen – Apologise – Solve – Thank; don't win the argument and lose the customer.",
    "Đa văn hóa: hiểu trước, phán xét không bao giờ": "Cross-cultural: understand first, never judge",
    "Tìm hiểu – quan sát – thích ứng; xác nhận thỏa thuận bằng văn bản để vượt rào cản ngôn ngữ.":
        "Research – observe – adapt; confirm agreements in writing to get past the language barrier.",
    "Trình bày cách báo cáo tin xấu với cấp trên qua một ví dụ cụ thể.": "Explain how to report bad news to a superior, using a concrete example.",
    "Vận dụng quy trình LAST để xử lý một tình huống phàn nàn tự chọn.": "Apply the LAST process to handle a complaint scenario of your choice.",
    "Nêu 5 điều nên làm và 5 điều nên tránh khi dự tiệc cùng đối tác kinh doanh.":
        "State 5 things to do and 5 things to avoid at a banquet with a business partner.",
    "Phân tích một khác biệt văn hóa Đông – Tây và cách thích ứng khi làm việc.":
        "Analyse one East–West cultural difference and how to adapt to it at work.",
    "CHUẨN BỊ CHO BUỔI SAU:  Chương 4 – Đàm phán trong kinh doanh. Mỗi nhóm nghĩ về lần “trả giá” gần nhất của mình (mua xe, thuê trọ…): điều gì khiến bạn thành công hoặc thất bại?":
        "PREPARE FOR NEXT SESSION:  Chapter 4 – Business negotiation. Each group should think of their most recent \"haggling\" experience (buying a vehicle, renting a room…): what made you succeed or fail?",
}

CH4 = {
    "Đàm phán": "Negotiation",
    "trong kinh doanh": "in business",
    "Nghệ thuật đạt thỏa thuận mà không đánh mất quan hệ — kỹ năng sinh lời trực tiếp nhất của người làm kinh tế.":
        "The art of reaching a deal without losing the relationship — the most directly profitable skill for anyone in business.",
    "Trình bày  —  khái niệm, đặc điểm và các kiểu đàm phán trong kinh doanh.":
        "Explain  —  the concept, features and styles of business negotiation.",
    "Mô tả  —  tiến trình đàm phán 5 giai đoạn và nhiệm vụ then chốt của từng giai đoạn.":
        "Describe  —  the 5-stage negotiation process and the key task of each stage.",
    "Vận dụng  —  các kỹ năng đàm phán cơ bản: chuẩn bị BATNA, đặt câu hỏi, nhượng bộ có điều kiện.":
        "Apply  —  basic negotiation skills: preparing a BATNA, questioning, and conditional concessions.",
    "Nhận diện  —  các chiêu trò thường gặp trên bàn đàm phán và cách ứng phó chuyên nghiệp.":
        "Recognise  —  common negotiation tactics and how to respond professionally.",
    "4.1  —  Khái niệm, đặc điểm và các kiểu đàm phán trong kinh doanh":
        "4.1  —  Concept, features and styles of business negotiation",
    "4.2  —  Tiến trình đàm phán qua năm giai đoạn, từ chuẩn bị đến sau đàm phán":
        "4.2  —  The five-stage negotiation process, from preparation to follow-up",
    "4.3  —  Các kỹ năng đàm phán và cách nhận diện chiêu trò thường gặp":
        "4.3  —  Negotiation skills and recognising common tactics",
    "MỤC 4.1": "SECTION 4.1",
    "Khái niệm và các kiểu đàm phán": "The concept and styles of negotiation",
    "Hiểu bản chất kép của đàm phán: vừa hợp tác vừa cạnh tranh.": "Understand negotiation's dual nature: both cooperative and competitive.",
    "Đàm phán là gì?": "What is negotiation?",
    "Khái niệm": "Concept",
    "Đàm phán là quá trình các bên vừa có lợi ích chung, vừa có lợi ích xung đột, cùng trao đổi – thuyết phục để đi đến một thỏa thuận mà các bên chấp nhận được.":
        "Negotiation is a process where parties with both shared and conflicting interests exchange and persuade to reach an agreement acceptable to all.",
    "Bản chất kép: hợp tác + cạnh tranh": "Dual nature: cooperation + competition",
    "Hợp tác để “chiếc bánh” tồn tại và lớn lên; cạnh tranh khi phân chia chiếc bánh — quên vế nào cũng thất bại.":
        "Cooperate so the \"pie\" exists and grows; compete when dividing it up — forgetting either half means failure.",
    "Ba nguồn sức mạnh trên bàn đàm phán": "Three sources of power at the table",
    "Thông tin (ai hiểu đối phương hơn) • Thời gian (ai ít bị ép tiến độ hơn) • Thế lực (ai có nhiều lựa chọn thay thế hơn).":
        "Information (who understands the other side better) • Time (who is under less pressure to finish) • Leverage (who has more alternatives).",
    "Đặc điểm của đàm phán trong kinh doanh": "Features of business negotiation",
    "Lấy lợi ích kinh tế làm trung tâm  —  mọi điều khoản cuối cùng đều quy về giá trị, chi phí, rủi ro của mỗi bên.":
        "Centred on economic interest  —  every final term boils down to each side's value, cost and risk.",
    "Các bên vừa phụ thuộc vừa độc lập  —  cần nhau để có thỏa thuận, nhưng mỗi bên luôn có phương án riêng của mình.":
        "Parties are both dependent and independent  —  they need each other to reach a deal, yet each always has its own alternative.",
    "Thỏa thuận phải được văn bản hóa  —  kết quả đàm phán chỉ an toàn khi thành hợp đồng đúng thể thức — cầu nối sang Chương 5.":
        "Agreements must be documented  —  a negotiated outcome is only safe once it becomes a properly formatted contract — the bridge to Chapter 5.",
    "Diễn ra trong giới hạn  —  thời gian, thẩm quyền, ngân sách; nhà đàm phán giỏi biết rõ giới hạn của mình và ước lượng giới hạn đối phương.":
        "Bounded by limits  —  time, authority, budget; a good negotiator knows their own limits and estimates the other side's.",
    "Chịu ảnh hưởng văn hóa và quan hệ  —  phong cách đàm phán Á – Âu khác nhau; thương vụ một lần khác quan hệ hợp tác lâu dài.":
        "Shaped by culture and relationships  —  negotiating styles differ between Asia and the West; a one-off deal differs from a long-term partnership.",
    "Các kiểu đàm phán": "Negotiation styles",
    "Đàm phán kiểu mềm": "Soft negotiation",
    "Coi đối tác như bạn, dễ nhượng bộ để giữ quan hệ — nhanh đạt thỏa thuận nhưng dễ chịu thiệt khi gặp đối thủ cứng.":
        "Treats the other side as a friend, concedes easily to preserve the relationship — reaches agreement fast but loses out against a hard negotiator.",
    "Đàm phán kiểu cứng": "Hard negotiation",
    "Coi đối tác như đối thủ, ép buộc, giữ lập trường đến cùng — có thể thắng một lần nhưng phá vỡ quan hệ, dễ bế tắc.":
        "Treats the other side as an adversary, pressures them, holds its position to the end — may win once but wrecks the relationship and often deadlocks.",
    "Đàm phán kiểu nguyên tắc (Harvard)": "Principled negotiation (Harvard)",
    "Tách con người khỏi vấn đề; tập trung vào lợi ích, không cố thủ lập trường; sáng tạo phương án cùng có lợi; dựa trên tiêu chí khách quan.":
        "Separate the people from the problem; focus on interests, not positions; invent options for mutual gain; insist on objective criteria.",
    "Phân bổ  ↔  Tích hợp": "Distributive  ↔  Integrative",
    "Phân bổ: chia chiếc bánh cố định (được – mất). Tích hợp: làm chiếc bánh lớn hơn bằng cách khai thác khác biệt về ưu tiên — hướng đến win-win.":
        "Distributive: dividing a fixed pie (win – lose). Integrative: growing the pie by exploiting differing priorities — aiming for a win-win.",
    "MỤC 4.2": "SECTION 4.2",
    "Tiến trình đàm phán năm giai đoạn": "The five-stage negotiation process",
    "Bảy mươi phần trăm kết quả được quyết định trước khi hai bên ngồi vào bàn.": "Seventy percent of the outcome is decided before both sides sit down at the table.",
    "Tiến trình đàm phán: 5 giai đoạn": "The negotiation process: 5 stages",
    "Giai đoạn chuẩn bị — vũ khí quan trọng nhất": "The preparation stage — your most important weapon",
    "ZOPA hẹp hay rộng phụ thuộc vào giới hạn thật của hai bên — chuẩn bị kỹ để biết mình đang ở đâu trên trục này.":
        "Whether the ZOPA is narrow or wide depends on both sides' real limits — prepare well to know where you stand on that range.",
    "CON SỐ PHẢI NHỚ": "THE NUMBER TO REMEMBER",
    "Vì sao chuẩn bị lại quan trọng đến thế?": "Why does preparation matter so much?",
    "KẾT QUẢ": "OF THE OUTCOME",
    "được quyết định ngay từ giai đoạn chuẩn bị, trước khi hai bên ngồi vào bàn": "is decided during preparation, before both sides even sit down",
    "MỨC MỤC TIÊU": "TARGET LEVELS",
    "Lý tưởng – Kỳ vọng – Tối thiểu: phải viết ra giấy trước khi đàm phán": "Ideal – Expected – Minimum: must be written down before negotiating",
    "BATNA": "BATNA",
    "Phương án thay thế tốt nhất — nguồn sức mạnh thật sự trên bàn đàm phán": "Best alternative to a negotiated agreement — the real source of power at the table",
    "Câu để đời của chương: “Không chuẩn bị chính là chuẩn bị để nhượng bộ.”":
        "The chapter's defining line: \"Failing to prepare is preparing to concede.\"",
    "Mở đầu và thương lượng": "Opening and bargaining",
    "Tạo không khí và thăm dò": "Setting the mood and probing",
    "Vài phút xã giao đúng mực; quan sát thái độ; đặt câu hỏi mở để đối phương bộc lộ nhu cầu và giới hạn trước khi mình ra giá.":
        "A few minutes of proper small talk; observe their attitude; ask open questions so the other side reveals their needs and limits before you name a price.",
    "Đưa đề nghị và mặc cả": "Making offers and bargaining",
    "Đề nghị đầu tiên có căn cứ (neo tâm lý); phản hồi đề nghị của đối phương bằng câu hỏi “dựa trên cơ sở nào?” thay vì đồng ý hay bác bỏ ngay.":
        "Make a well-grounded first offer (a psychological anchor); respond to their offer with \"on what basis?\" instead of accepting or rejecting immediately.",
    "Nhượng bộ có điều kiện": "Conditional concessions",
    "Không cho không bao giờ: “Nếu anh tăng số lượng lên 500, chúng tôi sẽ giảm 3%.” Nhượng bộ nhỏ dần để phát tín hiệu chạm giới hạn.":
        "Never give anything for free: \"If you raise the quantity to 500, we'll cut 3%.\" Make concessions progressively smaller to signal you're nearing your limit.",
    "Xử lý bế tắc": "Handling deadlock",
    "Tạm nghỉ; đổi người – đổi vấn đề – đổi cách tiếp cận; quay về lợi ích gốc; đưa tiêu chí khách quan (giá thị trường, quy định) làm trọng tài.":
        "Take a break; change the person, the issue, or the approach; return to the underlying interests; bring in objective criteria (market price, regulations) as an arbiter.",
    "Kết thúc và sau đàm phán": "Closing and after the negotiation",
    "Nhận biết thời điểm chốt": "Recognising when to close",
    "Đối phương hỏi chi tiết triển khai, điều khoản thanh toán, thời gian giao hàng — tín hiệu sẵn sàng; tóm tắt thỏa thuận và đề nghị xác nhận.":
        "The other side asks about implementation details, payment terms, delivery time — a readiness signal; summarise the deal and ask them to confirm.",
    "Văn bản hóa ngay": "Document it immediately",
    "Thỏa thuận miệng chưa phải kết thúc: lập biên bản ghi nhớ, soạn hợp đồng đủ điều khoản cơ bản — kỹ thuật soạn thảo học ở Chương 5 và phần thực hành.":
        "A verbal agreement isn't the end: draw up a memorandum, draft a contract with all the basic terms — the drafting technique covered in Chapter 5 and the practice sessions.",
    "Sau đàm phán": "After the negotiation",
    "Thực hiện đúng cam kết — uy tín cho lần đàm phán sau; giữ liên lạc với đối tác; họp nhóm rút kinh nghiệm: điều gì hiệu quả, điều gì cần làm khác.":
        "Deliver on commitments — it builds credibility for next time; stay in touch with the partner; hold a team debrief: what worked, what to do differently.",
    "MỤC 4.3": "SECTION 4.3",
    "Kỹ năng và chiêu trò trên bàn đàm phán": "Skills and tactics at the negotiating table",
    "Những gì cần rèn, và những gì cần nhận diện để không bị dẫn dắt.": "What to train, and what to recognise so you don't get played.",
    "Kỹ năng nền tảng trên bàn đàm phán": "Foundational negotiation skills",
    "Lắng nghe và đặt câu hỏi": "Listening and questioning",
    "Nghe nhiều hơn nói; hỏi mở để tìm lợi ích thật; im lặng đúng lúc — nhiều nhượng bộ xuất hiện chỉ vì đối phương không chịu được khoảng lặng.":
        "Listen more than you talk; ask open questions to find real interests; use silence at the right moment — many concessions happen simply because the other side can't bear the silence.",
    "Thuyết phục bằng lợi ích và bằng chứng": "Persuading with interests and evidence",
    "Nói bằng ngôn ngữ lợi ích của đối phương; kèm số liệu, tiền lệ, quy định khách quan thay vì tranh cãi cảm tính.":
        "Speak in terms of the other side's interests; back it up with data, precedent and objective rules instead of emotional arguing.",
    "Kiểm soát cảm xúc": "Controlling emotions",
    "Giữ bình tĩnh trước khiêu khích; tách con người khỏi vấn đề; tức giận là nhượng quyền kiểm soát cho đối phương.":
        "Stay calm under provocation; separate the person from the problem; getting angry hands control over to the other side.",
    "Làm việc theo êkíp": "Working as a team",
    "Phân vai trưởng đoàn – chuyên môn – ghi chép; thống nhất tín hiệu nội bộ; không bao giờ mâu thuẫn nội bộ trước mặt đối tác.":
        "Assign roles — lead, expert, note-taker; agree on internal signals; never disagree with each other in front of the other side.",
    "Nhận diện chiêu trò thường gặp": "Recognising common tactics",
    "Neo giá sốc  —  mở màn bằng đề nghị cao/thấp bất thường để kéo kỳ vọng của ta — ứng phó: bám vào tiêu chí khách quan, đừng vội điều chỉnh mục tiêu.":
        "Extreme anchoring  —  opening with an unusually high/low offer to shift our expectations — counter: stick to objective criteria, don't rush to adjust your target.",
    "Người tốt – kẻ xấu  —  một người gay gắt, một người “dễ thương” ra tay cứu vãn — nhận diện và chỉ đàm phán trên nội dung.":
        "Good cop, bad cop  —  one person is harsh, another \"nice\" one steps in to save the day — recognise it and negotiate on substance only.",
    "Thời hạn chót giả  —  “chỉ còn hôm nay” để ép quyết định vội — kiểm chứng thực hư, sẵn sàng rời bàn nếu có BATNA.":
        "Fake deadline  —  \"only today\" to force a hasty decision — verify if it's real, and be ready to walk away if you have a BATNA.",
    "Cắt lát salami  —  đòi thêm từng chút nhỏ sau khi đã thỏa thuận — gói toàn bộ điều khoản lại: “điểm này mở thì cả gói mở”.":
        "Salami slicing  —  asking for small extras one at a time after the deal is set — bundle every term back together: \"reopen this and the whole package reopens\".",
    "Đòi hỏi phút chót  —  thêm yêu cầu ngay trước khi ký — bình tĩnh định giá yêu cầu đó và đòi đối ứng tương xứng.":
        "Last-minute demand  —  adding a request right before signing — calmly price that request and demand something equivalent in return.",
    "Role-play: đàm phán mua thiết bị": "Role-play: negotiating an equipment purchase",
    "TÌNH HUỐNG (2 nhóm/cặp, 25 phút; mỗi bên nhận “hồ sơ mật” riêng của giảng viên)":
        "SCENARIO (2 groups/pairs, 25 minutes; each side gets its own \"confidential brief\" from the instructor)",
    "Công ty X cần mua 20 máy tính cho phòng làm việc mới, ngân sách tối đa 240 triệu, cần giao trong 3 tuần. Nhà cung cấp Y muốn bán giá tốt nhưng đang tồn kho model cũ và muốn ký hợp đồng bảo trì dài hạn. Hai bên chưa biết giới hạn của nhau.":
        "Company X needs to buy 20 computers for a new office, with a maximum budget of 240 million VND and a 3-week deadline. Supplier Y wants a good price but is holding old-model stock and wants a long-term maintenance contract. Neither side knows the other's limits yet.",
    "Mỗi bên 10 phút chuẩn bị: xác định mục tiêu 3 mức, BATNA và chiến lược nhượng bộ theo hồ sơ được phát.":
        "Each side gets 10 minutes to prepare: set 3-level targets, a BATNA and a concession strategy based on the brief provided.",
    "Đàm phán 10 phút trước lớp; các nhóm quan sát ghi lại: đề nghị neo, các nhượng bộ, chiêu trò (nếu có).":
        "Negotiate for 10 minutes in front of the class; observing groups note down: the anchor offer, the concessions, any tactics used.",
    "Cả lớp phân tích: thỏa thuận đạt được nằm ở đâu trong ZOPA? Bên nào chuẩn bị tốt hơn và vì sao?":
        "The class analyses: where does the deal fall within the ZOPA? Which side prepared better, and why?",
    "Đàm phán thắng từ trước khi ngồi vào bàn": "Negotiations are won before you sit down",
    "Mục tiêu 3 mức + BATNA + hiểu đối tác = 70% kết quả; không chuẩn bị là chuẩn bị để nhượng bộ.":
        "3-level targets + a BATNA + understanding the other side = 70% of the outcome; failing to prepare is preparing to concede.",
    "Đàm phán lợi ích, đừng cố thủ lập trường": "Negotiate interests, don't dig in on positions",
    "Hỏi “vì sao” để tìm lợi ích thật; nhượng bộ luôn kèm điều kiện; hướng tới thỏa thuận hai bên thực hiện được.":
        "Ask \"why\" to find the real interest; every concession comes with a condition; aim for a deal both sides can actually deliver.",
    "Thỏa thuận chỉ an toàn khi thành văn bản": "An agreement is only safe once it's in writing",
    "Chốt xong phải văn bản hóa thành hợp đồng đúng thể thức — đó là nội dung Chương 5: Soạn thảo và trình bày văn bản.":
        "Once closed, it must be documented as a properly formatted contract — that's the subject of Chapter 5: Drafting and presenting documents.",
    "Phân biệt đàm phán kiểu mềm, kiểu cứng và kiểu nguyên tắc; khi nào nên dùng kiểu nào?":
        "Distinguish soft, hard and principled negotiation; when should each be used?",
    "BATNA và ZOPA là gì? Xây dựng BATNA cho một tình huống thuê nhà trọ của sinh viên.":
        "What are BATNA and ZOPA? Build a BATNA for a student's room-rental scenario.",
    "Trình bày 5 giai đoạn của tiến trình đàm phán và nhiệm vụ chính của mỗi giai đoạn.":
        "Explain the 5 stages of the negotiation process and the main task of each stage.",
    "Nêu 3 chiêu trò thường gặp trong đàm phán và cách ứng phó.": "State 3 common negotiation tactics and how to counter them.",
    "CHUẨN BỊ CHO BUỔI SAU:  Chương 5 – Soạn thảo và trình bày văn bản: đọc trước Nghị định 30/2020/NĐ-CP (phần thể thức văn bản); phần thực hành sẽ soạn hợp đồng cho chính thương vụ vừa đàm phán hôm nay.":
        "PREPARE FOR NEXT SESSION:  Chapter 5 – Drafting and presenting documents: read Decree 30/2020/ND-CP in advance (the section on document formality); the practice session will draft a contract for the very deal negotiated today.",
}

CH5 = {
    "Je m'appelle Huong  •  GV. Đỗ Thùy Hương  •  EC6000TX":
        "Je m'appelle Huong  •  Instructor Đỗ Thùy Hương  •  EC6000TX",
    "EC6000TX – Kỹ năng giao tiếp và đàm phán trong kinh doanh (Hệ từ xa)":
        "EC6000TX – Business Communication and Negotiation Skills (Distance Learning)",
    "© Đỗ Thùy Hương, 2026 — Bài giảng biên soạn cho lớp giảng dạy từ xa. Vui lòng ghi nguồn khi sử dụng.":
        "© Đỗ Thùy Hương, 2026 — Lecture prepared for the distance-learning class. Please credit the source when reused.",
    "Soạn thảo và": "Drafting and",
    "trình bày văn bản": "presenting documents",
    "Từ lời nói sang chữ viết: biến mọi thỏa thuận thành văn bản đúng chuẩn, có giá trị pháp lý.":
        "From words to writing: turning every agreement into a properly formatted document with legal standing.",
    "Sau bài học này, sinh viên có thể": "After this lesson, students will be able to",
    "Trình bày  —  khái niệm văn bản và phân biệt được các nhóm văn bản trong tổ chức.":
        "Explain  —  the concept of a document and distinguish the groups of documents within an organisation.",
    "Áp dụng  —  yêu cầu về nội dung và 9 thành phần thể thức theo Nghị định 30/2020/NĐ-CP.":
        "Apply  —  the content requirements and the 9 formality components under Decree 30/2020/ND-CP.",
    "Soạn thảo  —  được quyết định, tờ trình, công văn, biên bản, báo cáo đúng bố cục.":
        "Draft  —  decisions, proposals, official letters, minutes and reports with the correct structure.",
    "Soạn thảo  —  được thư tín thương mại, báo giá và hợp đồng phục vụ giao dịch kinh doanh.":
        "Draft  —  commercial correspondence, quotations and contracts for business transactions.",
    "5.1  —  Khái niệm và phân loại văn bản": "5.1  —  The concept and classification of documents",
    "5.2  —  Các yêu cầu về nội dung và thể thức văn bản": "5.2  —  Content and formality requirements",
    "5.3  —  Soạn thảo văn bản hành chính thông dụng": "5.3  —  Drafting common administrative documents",
    "5.4  —  Soạn thảo văn bản thương mại": "5.4  —  Drafting commercial documents",
    "MỤC 5.1 – 5.2": "SECTIONS 5.1 – 5.2",
    "Khái niệm, phân loại và thể thức văn bản": "The concept, classification and formality of documents",
    "Nền tảng pháp lý và kỹ thuật: văn bản là gì và một trang văn bản đúng chuẩn trông thế nào.":
        "The legal and technical foundation: what a document is, and what a correctly formatted page looks like.",
    "Văn bản là gì?": "What is a document?",
    "Cách hiểu chung": "General understanding",
    "Văn bản là phương tiện ghi lại và truyền đạt thông tin bằng ngôn ngữ hay ký hiệu nhất định, hình thành trong hoạt động của cơ quan, tổ chức, doanh nghiệp.":
        "A document is a means of recording and conveying information through language or symbols, formed in the activities of agencies, organisations and businesses.",
    "Định nghĩa pháp lý — Nghị định 30/2020/NĐ-CP": "Legal definition — Decree 30/2020/ND-CP",
    "“Văn bản là thông tin thành văn được truyền đạt bằng ngôn ngữ hoặc ký hiệu, hình thành trong hoạt động của các cơ quan, tổ chức và được trình bày đúng thể thức, kỹ thuật theo quy định.”":
        "\"A document is written information conveyed in language or symbols, formed in the activities of agencies and organisations, and presented in the format and technique prescribed.\"",
    "Vai trò": "Role",
    "Phương tiện quản lý – điều hành • căn cứ pháp lý cho hoạt động • lưu trữ thông tin • thể hiện hình ảnh chuyên nghiệp của tổ chức.":
        "A tool for management and administration • a legal basis for activities • a record of information • a reflection of the organisation's professional image.",
    "Phân loại văn bản": "Classifying documents",
    "Văn bản quy phạm pháp luật": "Legal normative documents",
    "Chứa quy tắc xử sự chung, do cơ quan nhà nước có thẩm quyền ban hành: Luật, Nghị định, Thông tư. Doanh nghiệp không ban hành nhưng phải tuân thủ.":
        "Contain general rules of conduct, issued by a competent state agency: laws, decrees, circulars. Businesses don't issue these but must comply with them.",
    "Văn bản hành chính": "Administrative documents",
    "Loại gặp nhiều nhất — 29 loại theo NĐ 30/2020: quyết định cá biệt, công văn, thông báo, báo cáo, tờ trình, biên bản…":
        "The most common type — 29 types under Decree 30/2020: individual decisions, official letters, notices, reports, proposals, minutes…",
    "Văn bản chuyên ngành": "Specialised documents",
    "Hình thành trong nghiệp vụ chuyên môn: chứng từ kế toán, hồ sơ kỹ thuật, hồ sơ mời thầu.":
        "Formed within specific professional practice: accounting vouchers, technical dossiers, bidding dossiers.",
    "Văn bản thương mại": "Commercial documents",
    "Phục vụ giao dịch kinh doanh: thư tín thương mại, báo giá, đơn đặt hàng, hợp đồng — học kỹ ở mục 5.4.":
        "Serving business transactions: commercial correspondence, quotations, purchase orders, contracts — covered in detail in section 5.4.",
    "Bốn yêu cầu về nội dung": "Four content requirements",
    "Đúng mục đích, đúng thẩm quyền": "Correct purpose and authority",
    "Mỗi văn bản tập trung một chủ đề; ban hành đúng chức năng, nhiệm vụ của cơ quan, tổ chức.":
        "Each document focuses on one topic; it is issued within the agency's or organisation's proper function and mandate.",
    "Chính xác — khách quan": "Accurate and objective",
    "Thông tin, số liệu trung thực, có căn cứ, được kiểm chứng. Một con số sai có thể tạo hậu quả pháp lý lớn.":
        "Information and figures must be truthful, evidence-based and verified. A single wrong figure can create serious legal consequences.",
    "Rõ ràng — ngắn gọn — dễ hiểu": "Clear, concise and easy to understand",
    "Câu văn mạch lạc, không đa nghĩa; người nhận đọc một lần là hiểu đúng ý người soạn.":
        "Sentences should be coherent and unambiguous; the reader should understand the drafter's intent on a single read.",
    "Đúng pháp luật, đúng ngôn ngữ hành chính": "Legally correct and in proper administrative language",
    "Phù hợp quy định hiện hành; văn phong nghiêm túc, lịch sự, không dùng khẩu ngữ.":
        "Compliant with current regulations; a formal, polite style with no colloquial language.",
    "Chín thành phần thể thức — NĐ 30/2020/NĐ-CP": "The nine formality components — Decree 30/2020/ND-CP",
    "Kỹ thuật trình bày — những con số phải thuộc": "Presentation technique — the figures you must know by heart",
    "Khổ giấy và lề trang": "Paper size and margins",
    "Khổ A4 (210 × 297 mm). Lề trên, dưới: 20 – 25 mm • lề trái: 30 – 35 mm (để đóng gáy) • lề phải: 15 – 20 mm.":
        "A4 size (210 × 297 mm). Top and bottom margins: 20 – 25 mm • left margin: 30 – 35 mm (for binding) • right margin: 15 – 20 mm.",
    "Phông chữ": "Font",
    "Times New Roman, bộ mã Unicode, cỡ 13 – 14, màu đen.": "Times New Roman, Unicode encoding, size 13 – 14, black.",
    "Số trang": "Page numbers",
    "Đánh từ trang thứ hai, bằng chữ số Ả Rập, canh giữa theo lề trên.":
        "Numbered from page two, in Arabic numerals, centred within the top margin.",
    "Ngôn ngữ": "Language",
    "Tiếng Việt chuẩn mực; viết hoa, viết tắt đúng quy định; số liệu dùng chữ số Ả Rập.":
        "Standard Vietnamese; correct capitalisation and abbreviations; figures use Arabic numerals.",
    "CON SỐ PHẢI NHỚ": "THE NUMBERS TO REMEMBER",
    "Bốn con số của một trang văn bản đúng chuẩn": "Four numbers of a correctly formatted document page",
    "A4": "A4",
    "KHỔ GIẤY": "PAPER SIZE",
    "210 × 297 mm — không dùng khổ Letter": "210 × 297 mm — never use Letter size",
    "MILIMET": "MILLIMETRES",
    "Lề trái, rộng nhất để đóng gáy lưu trữ": "The left margin, the widest one, for binding and filing",
    "CỠ CHỮ": "FONT SIZE",
    "Times New Roman, bộ mã Unicode, màu đen": "Times New Roman, Unicode encoding, black",
    "THÀNH PHẦN": "COMPONENTS",
    "Số thành phần thể thức bắt buộc theo NĐ 30/2020": "The number of mandatory formality components under Decree 30/2020",
    "Lề trên và dưới 20 – 25 mm • lề phải 15 – 20 mm • số trang đánh từ trang thứ hai, canh giữa theo lề trên.":
        "Top and bottom margins 20 – 25 mm • right margin 15 – 20 mm • page numbers from page two, centred within the top margin.",
    "MỤC 5.3": "SECTION 5.3",
    "Soạn thảo văn bản hành chính thông dụng": "Drafting common administrative documents",
    "Năm loại văn bản dùng hằng ngày trong mọi cơ quan, tổ chức.": "Five document types used every day in every agency and organisation.",
    "Năm văn bản hành chính thông dụng": "Five common administrative documents",
    "Chuỗi văn bản của một thương vụ — sinh viên sẽ soạn lại đúng chuỗi này trong phần thực hành.":
        "The document chain for a business deal — students will draft this very chain in the practice sessions.",
    "Quyết định và Tờ trình": "Decisions and proposals",
    "Quyết định — khái niệm": "Decision — concept",
    "Văn bản do người có thẩm quyền ban hành để giải quyết một công việc cụ thể.":
        "A document issued by a competent person to resolve a specific matter.",
    "Quyết định — bố cục": "Decision — structure",
    "Phần căn cứ (pháp lý + thực tiễn, kết thúc bằng dấu chấm) → phần nội dung theo các Điều. Điều cuối ghi hiệu lực và đối tượng thi hành.":
        "The grounds section (legal + practical, ending with a full stop) → the content, organised into Articles. The final Article states the effective date and who must comply.",
    "Tờ trình — khái niệm": "Proposal — concept",
    "Văn bản đề xuất cấp có thẩm quyền phê duyệt chủ trương, phương án, đề án hoặc giải quyết công việc.":
        "A document proposing that a competent authority approve a policy, plan, project or course of action.",
    "Tờ trình — bố cục 3 phần": "Proposal — 3-part structure",
    "Mở đầu: lý do, sự cần thiết → Nội dung: phương án, lợi ích, tính khả thi → Kết thúc: kiến nghị phê duyệt. Đính kèm hồ sơ, dự toán.":
        "Opening: reason and necessity → Content: the plan, its benefits and feasibility → Closing: a request for approval. Dossiers and cost estimates are attached.",
    "Công văn, Biên bản và Báo cáo": "Official letters, minutes and reports",
    "Công văn": "Official letters",
    "Không có tên loại — chỉ có số, ký hiệu và trích yếu. Các loại: đề nghị, phúc đáp, đôn đốc, hướng dẫn, giải thích, mời họp. Mỗi công văn một chủ đề; kết thúc “Trân trọng./.”":
        "Has no type name — only a number, symbol and summary. Types include: request, response, follow-up, guidance, explanation, meeting invitation. Each letter covers one topic; it closes with \"Respectfully./.\"",
    "Biên bản": "Minutes",
    "Ghi tại chỗ, trung thực, khách quan. Kết cấu: thời gian – địa điểm → thành phần tham dự → diễn biến, ý kiến → kết luận → chữ ký các bên (yếu tố tạo giá trị pháp lý).":
        "Recorded on the spot, truthfully and objectively. Structure: time and place → attendees → proceedings and views → conclusion → the parties' signatures (what gives it legal validity).",
    "Báo cáo": "Reports",
    "Định kỳ • đột xuất • chuyên đề • sơ kết, tổng kết. Mạch 4 phần: đặc điểm tình hình → kết quả đạt được → hạn chế và nguyên nhân → phương hướng, kiến nghị.":
        "Periodic • ad-hoc • thematic • review/summary. A 4-part flow: the situation → results achieved → limitations and causes → direction and recommendations.",
    "MỤC 5.4": "SECTION 5.4",
    "Soạn thảo văn bản thương mại": "Drafting commercial documents",
    "Thư tín, báo giá và hợp đồng — bộ hồ sơ đưa một thương vụ đi đến đích.":
        "Correspondence, quotations and contracts — the set of documents that carries a deal to completion.",
    "Thư tín thương mại và báo giá": "Commercial correspondence and quotations",
    "Thư tín thương mại": "Commercial correspondence",
    "Thư hỏi hàng, chào hàng, đặt hàng, xác nhận, khiếu nại, cảm ơn. Kết cấu: mở đầu – nội dung – kết thúc.":
        "Inquiry, sales, order, confirmation, complaint and thank-you letters. Structure: opening – body – closing.",
    "Nguyên tắc 5C": "The 5C principle",
    "Clear (rõ) • Concise (gọn) • Correct (đúng) • Complete (đủ) • Courteous (lịch sự).":
        "Clear • Concise • Correct • Complete • Courteous.",
    "Email thương mại": "Commercial email",
    "Tiêu đề ngắn đúng nội dung; xưng hô phù hợp; chữ ký đầy đủ thông tin; phản hồi trong 24 giờ.":
        "A short subject line matching the content; an appropriate salutation; a signature with full details; reply within 24 hours.",
    "Báo giá": "Quotations",
    "Thông tin doanh nghiệp • mô tả hàng hóa • số lượng, đơn giá, thuế • điều kiện giao hàng, thanh toán • thời hạn hiệu lực (tránh tranh chấp khi giá thị trường biến động).":
        "Business details • description of goods • quantity, unit price, tax • delivery and payment terms • a validity period (to avoid disputes when market prices move).",
    "Hợp đồng — nghiệm thu — thanh lý": "Contract, acceptance and liquidation",
    "Hợp đồng thương mại": "Commercial contracts",
    "Căn cứ Bộ luật Dân sự 2015 và Luật Thương mại 2005. Điều khoản chính: đối tượng; giá và phương thức thanh toán; quyền – nghĩa vụ các bên; phạt vi phạm; giải quyết tranh chấp.":
        "Based on the 2015 Civil Code and the 2005 Commercial Law. Main clauses: subject matter; price and payment method; parties' rights and obligations; penalties; dispute resolution.",
    "Biên bản nghiệm thu": "Acceptance record",
    "Xác nhận khối lượng, chất lượng hàng hóa, dịch vụ đã thực hiện — căn cứ để thanh toán.":
        "Confirms the volume and quality of goods or services delivered — the basis for payment.",
    "Biên bản thanh lý hợp đồng": "Contract liquidation record",
    "Xác nhận hoàn thành nghĩa vụ, chấm dứt hiệu lực hợp đồng, quyết toán các quyền và nghĩa vụ còn lại.":
        "Confirms obligations are fulfilled, terminates the contract, and settles the remaining rights and obligations.",
    "Từ đàm phán đến hợp đồng": "From negotiation to contract",
    "TÌNH HUỐNG (nhóm 4–5 sinh viên, 20 phút)": "SCENARIO (groups of 4–5 students, 20 minutes)",
    "Tiếp nối thương vụ mua 20 máy tính đã đàm phán ở Chương 4: Phòng Hành chính Công ty X phải hoàn tất toàn bộ hồ sơ giấy tờ cho thương vụ, từ lúc đề xuất mua đến khi thanh toán xong.":
        "Continuing the 20-computer purchase negotiated in Chapter 4: Company X's Admin Office must complete the entire paper trail for the deal, from the initial purchase proposal through to final payment.",
    "Liệt kê đầy đủ chuỗi văn bản cần soạn theo đúng trình tự thời gian và cho biết ai ký từng văn bản.":
        "List the full chain of documents to draft, in chronological order, and state who signs each one.",
    "Chọn một văn bản trong chuỗi, phác thảo bố cục đầy đủ 9 thành phần thể thức lên giấy A4.":
        "Pick one document from the chain and sketch its full layout with all 9 formality components on A4 paper.",
    "Chỉ ra 3 lỗi thể thức thường gặp nhất mà nhóm dự đoán sinh viên hay mắc phải khi soạn văn bản này.":
        "Identify the 3 most common formatting mistakes your group expects students to make when drafting this document.",
    "Những điều cần nhớ": "Things to remember",
    "Thể thức là “giấy thông hành” của văn bản": "Formality is a document's \"passport\"",
    "Nội dung hay đến đâu mà sai thể thức thì văn bản vẫn bị trả lại — thuộc 9 thành phần và các con số trình bày.":
        "However good the content, a document with the wrong format still gets sent back — know the 9 components and the presentation figures.",
    "Mỗi loại văn bản có một bố cục riêng": "Each document type has its own structure",
    "Quyết định theo Điều; tờ trình 3 phần; công văn không tên loại; biên bản lập tại chỗ; báo cáo theo mạch 4 phần.":
        "Decisions follow Articles; proposals have 3 parts; official letters have no type name; minutes are recorded on the spot; reports follow a 4-part flow.",
    "Văn bản thương mại giữ uy tín doanh nghiệp": "Commercial documents protect a business's reputation",
    "Thư tín đạt 5C, báo giá có hiệu lực rõ, hợp đồng đủ điều khoản — bộ ba hợp đồng, nghiệm thu, thanh lý khép kín thương vụ.":
        "Correspondence following the 5Cs, a quotation with a clear validity period, a contract with full clauses — the contract/acceptance/liquidation trio closes out the deal.",
    "Trình bày khái niệm văn bản và các nhóm văn bản trong tổ chức.": "Explain the concept of a document and the groups of documents within an organisation.",
    "Nêu 9 thành phần thể thức văn bản hành chính và các quy định về lề trang, phông chữ.":
        "State the 9 formality components of an administrative document and the rules on margins and fonts.",
    "So sánh bố cục của quyết định, tờ trình và công văn.": "Compare the structure of a decision, a proposal and an official letter.",
    "Nêu nguyên tắc 5C và các điều khoản cơ bản của hợp đồng thương mại.": "State the 5C principle and the basic clauses of a commercial contract.",
    "CHUẨN BỊ CHO BUỔI SAU:  Phần thực hành tại phòng A0105 — Bài 1: Thể thức văn bản. Mang theo laptop, cài sẵn Microsoft Word và tải Nghị định 30/2020/NĐ-CP.":
        "PREPARE FOR NEXT SESSION:  Practice session in room A0105 — Lab 1: Document formality. Bring a laptop with Microsoft Word installed, and download Decree 30/2020/ND-CP in advance.",
    "Văn bản pháp lý bắt buộc": "Required legal document",
    "Nghị định 30/2020/NĐ-CP ngày 05/3/2020 của Chính phủ về công tác văn thư — hướng dẫn thể thức và kỹ thuật trình bày văn bản hành chính.":
        "Government Decree 30/2020/ND-CP dated 5 March 2020 on records and archival work — guiding the format and presentation technique for administrative documents.",
    "Slide, biểu mẫu văn bản, tình huống và bài tập do GV. Đỗ Thùy Hương biên soạn; cung cấp sau mỗi buổi học.":
        "Slides, document templates, case scenarios and exercises prepared by Instructor Đỗ Thùy Hương; shared after each session.",
}

DECKS = [
    ("slides/01-tong-quan-giao-tiep.pptx", CH1),
    ("slides/02-ky-nang-chuyen-nghiep.pptx", CH2),
    ("slides/03-tinh-huong-dac-thu.pptx", CH3),
    ("slides/04-dam-phan.pptx", CH4),
    ("slides/05-soan-thao-van-ban.pptx", CH5),
]


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # GROUP
            yield from walk(sh.shapes)
        else:
            yield sh


def dich_doan(p, dich):
    """Thay chữ ở run đầu của đoạn văn bằng bản dịch, xoá chữ các run thừa."""
    if not p.runs:
        return
    p.runs[0].text = dich
    for r in p.runs[1:]:
        r.text = ""


def dich_deck(duong_dan, bang_rieng):
    bang = {**CHUNG, **bang_rieng}
    prs = Presentation(duong_dan)
    thieu = []
    for s in prs.slides:
        for sh in walk(s.shapes):
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                goc = "".join(r.text for r in p.runs)
                if not goc.strip() or GIU_NGUYEN.match(goc):
                    continue
                if goc in bang:
                    dich_doan(p, bang[goc])
                else:
                    thieu.append(goc)
    if thieu:
        print(f"THIẾU BẢN DỊCH trong {duong_dan}:")
        for t in thieu:
            print("  ", repr(t))
        return False
    ra = duong_dan[:-5] + ".en.pptx"
    prs.save(ra)
    print(f"OK: {ra}")
    return True


if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    ok = True
    for duong_dan, bang in DECKS:
        if not os.path.exists(duong_dan):
            print("thiếu file:", duong_dan)
            ok = False
            continue
        ok = dich_deck(duong_dan, bang) and ok
    sys.exit(0 if ok else 1)
